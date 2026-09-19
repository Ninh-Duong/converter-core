import os
import re
import io
from pathlib import Path
from core.bootstrap import TESSERACT_CMD, get_gemini_api_key
from core.i18n import t

def is_scanned_pdf(pdf_path: Path) -> bool:
    """Check if a PDF file has an embedded digital text layer or is scanned/image-only."""
    import pymupdf
    doc = pymupdf.open(str(pdf_path))
    total_chars = sum(len(page.get_text().strip()) for page in doc)
    doc.close()
    return total_chars < 50

# -------------------------------------------------------------
# 1. Vision AI Engine (Gemini Flash)
# -------------------------------------------------------------
def process_page_with_vision_ai(img_bytes: bytes, api_key: str) -> str:
    """Transcribe and structure document image into Markdown using Gemini Flash Vision."""
    from google import genai
    from google.genai import types

    client = genai.Client(api_key=api_key)
    prompt = (
        "Convert this document page into clean standard Markdown format:\n"
        "- Preserve 100% accurate text, accents, and typography.\n"
        "- Use #, ##, ### for titles and headings.\n"
        "- Do not break sentences across arbitrary newlines; join paragraphs smoothly.\n"
        "- Represent tables as clean Markdown tables (| Col 1 | Col 2 |).\n"
        "- Return raw Markdown only, no code fences."
    )

    response = client.models.generate_content(
        model="gemini-2.5-flash",
        contents=[
            types.Part.from_bytes(data=img_bytes, mime_type="image/png"),
            prompt
        ]
    )
    return response.text

def render_markdown_to_docx(md_text: str, doc):
    """Render structured Markdown text into Microsoft Word DOCX elements."""
    from docx.shared import Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    lines = md_text.split("\n")
    in_table = False
    table_rows = []

    def flush_table():
        nonlocal in_table, table_rows
        if not table_rows:
            return
        parsed_rows = []
        for r in table_rows:
            if re.match(r"^\s*\|?\s*:?-+:?\s*(\|?\s*:?-+:?\s*)+\|?\s*$", r):
                continue
            cols = [c.strip() for c in r.strip("|").split("|")]
            parsed_rows.append(cols)
        
        if parsed_rows:
            max_cols = max(len(r) for r in parsed_rows)
            table = doc.add_table(rows=len(parsed_rows), cols=max_cols)
            table.style = "Table Grid"
            for r_idx, row in enumerate(parsed_rows):
                for c_idx, cell_text in enumerate(row):
                    if c_idx < max_cols:
                        cell = table.cell(r_idx, c_idx)
                        cell.text = cell_text
                        if r_idx == 0:
                            for p in cell.paragraphs:
                                for run in p.runs:
                                    run.bold = True
        table_rows = []
        in_table = False

    for line in lines:
        stripped = line.strip()
        if not stripped:
            if in_table:
                flush_table()
            continue

        if stripped.startswith("|") and stripped.endswith("|"):
            in_table = True
            table_rows.append(stripped)
            continue
        elif in_table:
            flush_table()

        if stripped.startswith("# "):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(stripped[2:].strip())
            run.bold = True
            run.font.size = Pt(14)
        elif stripped.startswith("## "):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(stripped[3:].strip())
            run.bold = True
            run.font.size = Pt(12)
        elif stripped.startswith("### "):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(stripped[4:].strip())
            run.bold = True
            run.font.size = Pt(11)
        elif stripped.startswith(("- ", "* ")):
            doc.add_paragraph(stripped[2:].strip(), style="List Bullet")
        else:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.line_spacing = 1.15
            parts = re.split(r"(\*\*.*?\*\*)", stripped)
            for part in parts:
                if part.startswith("**") and part.endswith("**"):
                    run = p.add_run(part[2:-2])
                    run.bold = True
                else:
                    p.add_run(part)

    if in_table:
        flush_table()

def pdf_vision_ai_to_docx(src: Path, dst: Path, api_key: str):
    import pymupdf, docx
    from docx.shared import Inches

    print(t("running_vision_ai"))
    pdf_in = pymupdf.open(str(src))
    doc_out = docx.Document()

    for sec in doc_out.sections:
        sec.top_margin = Inches(1)
        sec.bottom_margin = Inches(1)
        sec.left_margin = Inches(1)
        sec.right_margin = Inches(1)

def get_page_image_slices(page, pdf_in) -> list[bytes]:
    """Extract page image(s), automatically splitting 2-up landscape scans into portrait sub-pages."""
    from PIL import Image
    images = page.get_images()
    raw = pdf_in.extract_image(images[0][0])["image"] if images else page.get_pixmap(dpi=200).tobytes("png")
    img = Image.open(io.BytesIO(raw))
    w, h = img.size
    if w > h * 1.25:
        slices = []
        for box in ((0, 0, w // 2, h), (w // 2, 0, w, h)):
            buf = io.BytesIO()
            img.crop(box).save(buf, format="PNG")
            slices.append(buf.getvalue())
        return slices
    return [raw]

def save_docx_safely(doc_out, dst: Path) -> Path:
    """Save Word document, safely falling back to <stem>_fixed.docx if target is open/locked in Word."""
    try:
        doc_out.save(str(dst))
        return dst
    except PermissionError:
        fallback = dst.with_name(f"{dst.stem}_fixed{dst.suffix}")
        doc_out.save(str(fallback))
        print(f"\n[!] Warning: '{dst.name}' is open/locked. Saved to '{fallback.name}' instead.")
        return fallback

def pdf_vision_ai_to_docx(src: Path, dst: Path, api_key: str):
    import pymupdf, docx
    from docx.shared import Inches

    print(t("running_vision_ai"))
    pdf_in = pymupdf.open(str(src))
    doc_out = docx.Document()

    for sec in doc_out.sections:
        sec.top_margin = Inches(1)
        sec.bottom_margin = Inches(1)
        sec.left_margin = Inches(1)
        sec.right_margin = Inches(1)

    page_slices = []
    for page in pdf_in:
        page_slices.extend(get_page_image_slices(page, pdf_in))

    for idx, img_bytes in enumerate(page_slices, 1):
        print(t("page_progress_ai", page=idx, total=len(page_slices)))
        md_text = process_page_with_vision_ai(img_bytes, api_key)
        render_markdown_to_docx(md_text, doc_out)
        if idx < len(page_slices):
            doc_out.add_page_break()

    pdf_in.close()
    save_docx_safely(doc_out, dst)

# -------------------------------------------------------------
# 2. Local Tesseract OCR Engine (Offline Fallback)
# -------------------------------------------------------------
def extract_page_layout_tesseract(img_bytes: bytes, tess_cmd: str):
    import pytesseract
    from PIL import Image
    from pytesseract import Output

    pytesseract.pytesseract.tesseract_cmd = tess_cmd
    img = Image.open(io.BytesIO(img_bytes))
    _, h_total = img.size
    data = pytesseract.image_to_data(img, lang="vie", config="--psm 1", output_type=Output.DICT)

    raw_lines = []
    cur_key, cur_line = None, None
    for i in range(len(data["text"])):
        w = data["text"][i].strip()
        if not w:
            continue
        key = (data["block_num"][i], data["par_num"][i], data["line_num"][i])
        if key != cur_key:
            cur_key = key
            cur_line = {
                "words": [w],
                "left": data["left"][i],
                "right": data["left"][i] + data["width"][i],
                "top": data["top"][i],
                "bottom": data["top"][i] + data["height"][i],
            }
            raw_lines.append(cur_line)
        else:
            cur_line["words"].append(w)
            cur_line["right"] = max(cur_line["right"], data["left"][i] + data["width"][i])
            cur_line["bottom"] = max(cur_line["bottom"], data["top"][i] + data["height"][i])

    # Sort lines vertically to prevent out-of-order column/block jumps
    raw_lines.sort(key=lambda l: (round(l["top"] / 10) * 10, l["left"]))

    # Merge line fragments on the same baseline (e.g. margin overflows like "bảo vệ")
    merged = []
    for l in raw_lines:
        if merged and abs(merged[-1]["top"] - l["top"]) <= 8:
            merged[-1]["words"].extend(l["words"])
            merged[-1]["right"] = max(merged[-1]["right"], l["right"])
            merged[-1]["bottom"] = max(merged[-1]["bottom"], l["bottom"])
        else:
            merged.append(l)

    max_w = max((l["right"] - l["left"] for l in merged), default=1)
    page_elements = []
    cur_body = []

    for idx, l in enumerate(merged):
        text = " ".join(l["words"]).strip()
        if not text:
            continue

        # Ignore solitary page numbers in footer
        if text.isdigit() and len(text) <= 3 and l["top"] > h_total * 0.85:
            continue

        lw = l["right"] - l["left"]
        is_short = lw < (max_w * 0.82)
        is_title = text.isupper() or text.startswith(("CHƯƠNG", "Điều "))
        is_heading = bool(re.match(r"^(\d+(\.\d+)*|Mục|\>|\-|\•)\b", text))

        if is_title or is_heading:
            if cur_body:
                page_elements.append(("body", " ".join(cur_body)))
                cur_body = []
            page_elements.append(("title" if is_title else "heading", text))
            continue

        cur_body.append(text)

        next_is_heading = False
        if idx + 1 < len(merged):
            nxt = " ".join(merged[idx + 1]["words"]).strip()
            if nxt.isupper() or nxt.startswith(("CHƯƠNG", "Điều ")) or re.match(r"^(\d+(\.\d+)*|Mục|\>|\-|\•)\b", nxt):
                next_is_heading = True

        if is_short or next_is_heading:
            page_elements.append(("body", " ".join(cur_body)))
            cur_body = []

    if cur_body:
        page_elements.append(("body", " ".join(cur_body)))

    return page_elements

def render_tesseract_elements_to_docx(doc_out, elements):
    from docx.shared import Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    for elem_type, text in elements:
        text = text.strip()
        if not text or text in ["]", "[", "|", "~"]:
            continue

        p = doc_out.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.line_spacing = 1.15

        if elem_type == "title":
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(text)
            run.bold = True
            run.font.size = Pt(13)
        elif elem_type == "heading":
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(text)
            run.bold = True
            run.font.size = Pt(12)
        else:
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            run = p.add_run(text)
            run.font.size = Pt(12)

def pdf_tesseract_to_docx(src: Path, dst: Path):
    import pymupdf, docx
    from docx.shared import Inches

    if not TESSERACT_CMD:
        print(t("tesseract_not_found"))
        return

    print(t("running_tesseract"))
    pdf_in = pymupdf.open(str(src))
    doc_out = docx.Document()

    for sec in doc_out.sections:
        sec.top_margin = Inches(1)
        sec.bottom_margin = Inches(1)
        sec.left_margin = Inches(1)
        sec.right_margin = Inches(1)

    page_slices = []
    for page in pdf_in:
        page_slices.extend(get_page_image_slices(page, pdf_in))

    for idx, img_bytes in enumerate(page_slices, 1):
        print(t("page_progress_tess", page=idx, total=len(page_slices)))
        elements = extract_page_layout_tesseract(img_bytes, TESSERACT_CMD)
        render_tesseract_elements_to_docx(doc_out, elements)
        if idx < len(page_slices):
            doc_out.add_page_break()

    pdf_in.close()
    save_docx_safely(doc_out, dst)

# -------------------------------------------------------------
# 3. Fast Vector PDF Engine (pdf2docx)
# -------------------------------------------------------------
def pdf_layout_to_docx(src: Path, dst: Path):
    from pdf2docx import Converter
    cv = Converter(str(src))
    cv.convert(str(dst))
    cv.close()

# -------------------------------------------------------------
# 4. Hybrid Routers & Office Engines
# -------------------------------------------------------------
def convert_pdf_hybrid(src: Path, dst: Path):
    if not is_scanned_pdf(src):
        print(t("detected_vector_pdf"))
        pdf_layout_to_docx(src, dst)
        return

    print(t("detected_scanned_pdf"))
    api_key = get_gemini_api_key()
    if api_key:
        pdf_vision_ai_to_docx(src, dst, api_key)
    else:
        print(t("hint_no_api_key"))
        pdf_tesseract_to_docx(src, dst)

def image_to_docx_hybrid(src: Path, dst: Path):
    with open(src, "rb") as f:
        img_bytes = f.read()

    api_key = get_gemini_api_key()
    if api_key:
        import docx
        from docx.shared import Inches
        print(t("running_vision_ai"))
        doc_out = docx.Document()
        for sec in doc_out.sections:
            sec.top_margin = Inches(1)
            sec.bottom_margin = Inches(1)
            sec.left_margin = Inches(1)
            sec.right_margin = Inches(1)
        md_text = process_page_with_vision_ai(img_bytes, api_key)
        render_markdown_to_docx(md_text, doc_out)
        doc_out.save(str(dst))
    else:
        import docx
        from docx.shared import Inches
        print(t("running_tesseract"))
        doc_out = docx.Document()
        for sec in doc_out.sections:
            sec.top_margin = Inches(1)
            sec.bottom_margin = Inches(1)
            sec.left_margin = Inches(1)
            sec.right_margin = Inches(1)
        elements = extract_page_layout_tesseract(img_bytes, TESSERACT_CMD)
        render_tesseract_elements_to_docx(doc_out, elements)
        doc_out.save(str(dst))

def docx_to_pdf(src: Path, dst: Path):
    from docx2pdf import convert as d2p
    d2p(str(src), str(dst))

def get_converter_options(ext: str):
    """Return conversion options available for a file extension."""
    ext = ext.lower()
    if ext == ".pdf":
        return [
            ("mode_pdf_hybrid", "docx", convert_pdf_hybrid),
            ("mode_pdf_tesseract", "docx", pdf_tesseract_to_docx),
        ]
    elif ext in [".docx", ".doc"]:
        return [
            ("mode_docx_pdf", "pdf", docx_to_pdf),
        ]
    elif ext in [".png", ".jpg", ".jpeg"]:
        return [
            ("mode_image_docx", "docx", image_to_docx_hybrid),
        ]
    return []
